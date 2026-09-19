# -*- coding: utf-8 -*-
import io, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

src = io.open('shahruhiya/SCENARIY.md', encoding='utf-8').read().split('\n')

doc = Document()
st = doc.styles['Normal']
st.font.name = 'DejaVu Serif'
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(4)
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)

INLINE = re.compile(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)')

def runs(p, text, base_italic=False, base_bold=False, mono=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i = base_bold, base_italic
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, i = part[1:-1], True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            part = part[1:-1]
            mono = True
        r = p.add_run(part)
        r.bold, r.italic = b, i
        r.font.name = 'DejaVu Sans Mono' if mono else 'DejaVu Serif'
    return p

def heading(text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if level == 1:
        pf.space_before, pf.space_after = Pt(18), Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(18)
    elif level == 2:
        pf.space_before, pf.space_after = Pt(16), Pt(5)
        pf.keep_with_next = True
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(12)
    else:
        pf.space_before, pf.space_after = Pt(12), Pt(3)
        pf.keep_with_next = True
        r = p.add_run(text); r.bold = True; r.font.size = Pt(11)
    r.font.name = 'DejaVu Serif'

def flush_table(rows):
    rows = [r for r in rows if not re.match(r'^\s*\|[\s:\-|]+\|\s*$', r)]
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    if not cells:
        return
    n = max(len(c) for c in cells)
    t = doc.add_table(rows=0, cols=n)
    t.style = 'Table Grid'
    for ri, row in enumerate(cells):
        row = row + [''] * (n - len(row))
        tr = t.add_row().cells
        for ci, val in enumerate(row):
            para = tr[ci].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            runs(para, val, base_bold=(ri == 0))
            for r in para.runs:
                r.font.size = Pt(8)
    doc.add_paragraph()

i = 0
tbuf = []
while i < len(src):
    line = src[i].rstrip()
    if line.startswith('|'):
        tbuf.append(line); i += 1; continue
    if tbuf:
        flush_table(tbuf); tbuf = []
    if not line.strip():
        i += 1; continue
    if line.startswith('###'):
        heading(line.lstrip('#').strip(), 3)
    elif line.startswith('##'):
        heading(line.lstrip('#').strip(), 2)
    elif line.startswith('#'):
        heading(line.lstrip('#').strip(), 1)
    elif re.match(r'^-{3,}$', line.strip()):
        pass
    elif line.startswith('>'):
        body = line.lstrip('>').strip()
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        speaker = body.startswith('**')
        pf.left_indent = Cm(5.0) if speaker else Cm(6.2)
        if speaker:
            pf.space_before = Pt(8)
        runs(p, body, mono=True)
        for r in p.runs:
            r.font.size = Pt(10.5)
            if r.italic:
                r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
                r.font.size = Pt(9)
    elif re.match(r'^\s*[-*]\s+', line):
        p = doc.add_paragraph(style='List Bullet')
        runs(p, re.sub(r'^\s*[-*]\s+', '', line))
    elif re.match(r'^\s*\d+\.\s+', line):
        p = doc.add_paragraph(style='List Number')
        runs(p, re.sub(r'^\s*\d+\.\s+', '', line))
    else:
        runs(doc.add_paragraph(), line)
    i += 1
if tbuf:
    flush_table(tbuf)

out = 'shahruhiya/export/SHAHRUHIYA-scenariy.docx'
doc.save(out)
print('docx ok', out)
